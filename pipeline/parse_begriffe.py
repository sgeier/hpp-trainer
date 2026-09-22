"""Export the vocabulary quiz from Lernstuff/Begriffe/HPP Fachbegriffe.xlsx (391 terms with quiz answer + distractor)."""
import openpyxl, json, os

SRC = "/Users/sgeier/Projects/Phine-Lernen/HPP 2026/Lernstuff/Begriffe/HPP Fachbegriffe.xlsx"
OUT = os.path.join(os.path.dirname(__file__), "out")


def main():
    ws = openpyxl.load_workbook(SRC, read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c else "" for c in rows[0]]
    print("header:", header)
    # header row is offset: first column has no header ("Begriff"), so map by position
    out = []
    for r in rows[1:]:
        cells = [("" if c is None else str(c).strip()) for c in r]
        if not cells or not cells[0]:
            continue
        term, erkl = cells[0], cells[1] if len(cells) > 1 else ""
        rec = {"term": term, "explanation": erkl}
        for i, h in enumerate(header):
            if i < len(cells) and h:
                rec[h] = cells[i]
        out.append(rec)
    keys = collections = {}
    print("terms:", len(out), "with quiz answer:", sum(1 for r in out if r.get("Quiz: richtige Antwort")),
          "with distractor:", sum(1 for r in out if r.get("Quiz: Distraktor")))
    print("sample:", json.dumps(out[5], ensure_ascii=False)[:400])
    json.dump(out, open(os.path.join(OUT, "begriffe.json"), "w"), ensure_ascii=False, indent=1)


if __name__ == "__main__":
    main()


# ---------------------------------------------------------------------------
# Flashcard decks from the other sheets (term + explanation only, no distractors)
CARD_SRC = "/Users/sgeier/Projects/Phine-Lernen/HPP 2026/Lernstuff/Begriffe/"


def _rows(wb, name):
    ws = wb[name]
    return [[("" if c is None else str(c).strip()) for c in r] for r in ws.iter_rows(values_only=True)]


def cards():
    out = []
    wb = openpyxl.load_workbook(CARD_SRC + "Vokabeln und Begriffe.xlsx", read_only=True, data_only=True)
    merged = {}
    for r in _rows(wb, "allgemein")[1:]:
        if r[0] and r[1]:
            merged.setdefault(r[0], []).append(r[1])
    out += [{"deck": "Allgemein", "term": t, "text": " · ".join(dict.fromkeys(v))} for t, v in merged.items()]
    out += [{"deck": "Medizinisch", "term": r[1], "text": r[2]} for r in _rows(wb, "medizinisch")[1:] if len(r) > 2 and r[1] and r[2]]
    for r in _rows(wb, "Abkürzungen")[1:]:
        r += [""] * (9 - len(r))
        if r[1] and r[2]:
            out.append({"deck": "Abkürzungen", "term": r[1], "text": r[2], "note": r[3]})
        if r[6] and r[7]:
            out.append({"deck": "Wortbausteine", "term": r[6], "text": r[7], "note": r[8]})
    out += [{"deck": "Abwehrmechanismen", "term": r[0], "text": r[1]} for r in _rows(wb, "Abwehrmechanismen")[1:] if r[0] and r[1]]
    wb2 = openpyxl.load_workbook(CARD_SRC + "Fremdwörter_von_Donatella.xlsx", read_only=True, data_only=True)
    out += [{"deck": "Fremdwörter", "term": r[0], "text": r[1], "note": r[2], "category": r[3]} for r in _rows(wb2, "Tabelle1")[1:] if r[0] and r[1]]
    for c in out:
        c["note"] = c.get("note", "") or ""; c["category"] = c.get("category", "") or ""
    import collections
    print("cards:", dict(collections.Counter(c["deck"] for c in out)))
    json.dump(out, open(os.path.join(OUT, "cards.json"), "w"), ensure_ascii=False, indent=1)


if __name__ == "__main__":
    cards()
